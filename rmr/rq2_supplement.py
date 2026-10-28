"""Steps 8-10: supplementary extraction of RQ2 (form of the solution and coding activity).

Added during the second-round revision of the article. The metaprotocol's extraction form for
RQ2 ("How can these solutions be used?") asks for the form of the solution (library, plugin,
application, web service, IDE integration, command-line tool) and the context in which it is
applied, but the first execution recorded only the usage context (``RQb``). These steps extract
the two missing facets for the INCLUDED studies with the same procedure as step 4: an agent
answers from the full text, and one reviewer validates every answer against the text.

- **Step 8** (per origin): for each study the humans included at step 4 (``human_decision`` =
  include), run the agent over its full text and store ``form`` and ``activity`` in
  ``data/screen/<MODEL_EXPERIMENT>/<origin>/step-8-rq2-supplement.json``. Resumable: a record
  already answered is skipped. The step-4 file is only read, never changed.
- **Step 9** (all origins): export ``data/review/<MODEL_EXPERIMENT>/step-8-rq2-review.xlsx``,
  one row per study, with the agent's answers and the reviewer columns, prefilled with the
  agent's values. The reviewer corrects a value only when the text contradicts it, and sets
  ``Checked`` to Yes on every row once it is reviewed.
- **Step 10** (all origins): read the reviewed sheet back and store each validated value on the
  step-8 records as ``human_value`` next to the model's ``value`` (only where they differ), the
  same convention as the step-4 answers. Rows not marked as checked are reported, not applied.

The usage context (``RQb``) is not re-extracted: it stays as published.
"""

import json
from datetime import datetime, timezone
from typing import Literal

from agents import Agent, set_tracing_disabled
from pydantic import BaseModel

from rmr.llm import chat_model, model_settings
from rmr.paths import (PROJECT_ROOT, ensure_parent, full_path, review_dir,
                       step_output_path)
from rmr.screening.fulltext import FULLTEXT_MAX_CHARS, _run_agent

set_tracing_disabled(True)

PROMPT = PROJECT_ROOT / "prompts" / "step8-rq2-form-activity.md"
MAX_TOKENS = 35662  # same output budget as the step-4 answers (two facets fit easily)
ORIGINS = ["scopus", "google", "github", "hf"]
ORIGIN_BY_PREFIX = {"SC": "scopus", "GO": "google", "GH": "github", "HF": "hf"}
PENDING, ANSWERED = "pending", "answered"

# The closed sets, defined before the run (see the prompt for each definition). They feed both
# the structured output type, which restricts the model to these values, and the dropdowns of
# the review sheet, so the agent and the reviewer choose from the same lists.
FORMS = ("Tool", "Library or framework", "Model", "Agent", "Method or technique",
         "Dataset or benchmark", "Other")
TOOL_SUBTYPES = ("Application or platform", "Plugin or IDE integration", "Web service or API",
                 "Command-line tool")
ACTIVITIES = ("Code generation", "Code completion", "Code evolution",
              "Refactoring and restructuring", "Debugging and repair",
              "Static analysis and code review", "Construction testing", "Other")


class FormAnswer(BaseModel):
    note: str
    value: Literal[FORMS]
    tool_subtype: Literal[TOOL_SUBTYPES + ("",)] = ""  # only when value is "Tool"
    other: str = ""  # the form's name when value is "Other"
    evidence: str


class ActivityAnswer(BaseModel):
    note: str
    value: Literal[ACTIVITIES]
    other: str = ""  # the activity's name when value is "Other"
    evidence: str


class RQ2Supplement(BaseModel):
    form: FormAnswer
    activity: ActivityAnswer


# --- Step 8: extraction ---

def _included(origin: str) -> list[dict]:
    """Step-4 records the humans included: the catalogue this step classifies."""
    data = json.loads(step_output_path(origin, 4).read_text(encoding="utf-8"))
    return [r for r in data.get("records", []) if r.get("human_decision") == "include"]


def _new_record(source: dict) -> dict:
    answers = source.get("answers") or {}
    return {
        "id": source["id"],
        "status": PENDING,
        "data": {"title": source.get("data", {}).get("title", ""),
                 "solution_name": answers.get("solution_name", "")},
        "rq2": None,
    }


def _load(origin: str) -> tuple[dict, dict]:
    path = step_output_path(origin, 8)
    if not path.exists():
        return {}, {}
    data = json.loads(path.read_text(encoding="utf-8"))
    return data, {r["id"]: r for r in data.get("records", [])}


def _save(origin: str, records: list[dict], model: str | None) -> None:
    output = {
        "origin": origin,
        "step": 8,
        "model": model,
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "count": len(records),
        "by_status": {st: sum(1 for r in records if r["status"] == st) for st in (PENDING, ANSWERED)},
        "records": records,
    }
    path = step_output_path(origin, 8)
    ensure_parent(path)
    path.write_text(json.dumps(output, indent=2, ensure_ascii=False), encoding="utf-8")


def step8_rq2_supplement(origin: str) -> dict:
    included = _included(origin)
    data, existing = _load(origin)
    records = [existing.get(s["id"]) or _new_record(s) for s in included]
    state = {"model": data.get("model")}
    _save(origin, records, state["model"])  # first action: the manifest is the source of truth
    print(f"[{origin}] step 8: {len(records)} included studies")

    llm, model_name = chat_model()
    state["model"] = model_name
    agent = Agent(
        name="RQ2 form and activity classifier",
        instructions=PROMPT.read_text(encoding="utf-8"),
        model=llm,
        model_settings=model_settings(max_tokens=MAX_TOKENS),
        output_type=RQ2Supplement,
    )

    for record in records:
        rid = record["id"]
        if record["status"] == ANSWERED:
            continue
        path = full_path(origin, rid)
        text = path.read_text(encoding="utf-8").strip() if path.exists() else ""
        if not text:
            print(f"[{origin}] rq2 {rid}: no full text at {path} -> pending")
            continue
        message = (f"SOLUTION NAME: {record['data']['solution_name']}\n\n"
                   f"FULL TEXT:\n{text[:FULLTEXT_MAX_CHARS]}")
        result = _run_agent(agent, message, origin, rid, "rq2")
        if result is None:  # a bad response must not abort the batch; retry on the next run
            _save(origin, records, state["model"])
            continue
        record["rq2"] = result.model_dump()
        record["status"] = ANSWERED
        print(f"[{origin}] rq2 {rid}: form={result.form.value} | activity={result.activity.value}")
        _save(origin, records, state["model"])

    return {"origin": origin, "count": len(records),
            "answered": sum(1 for r in records if r["status"] == ANSWERED)}


# --- Step 9: review sheet ---

REVIEW_SHEET = "step-8-rq2-review.xlsx"
CHECKED = "Checked"
# Reviewer columns, prefilled with the agent's values; step 10 reads them back by name.
FORM_COL, SUBTYPE_COL, ACTIVITY_COL = "Form", "Tool subtype", "Activity"
FORM_OTHER_COL, ACTIVITY_OTHER_COL, COMMENT_COL = "Form other", "Activity other", "Comment"
HEADER = ["ID", "Title", "Solution name", "Usage context (RQ2)", "Form (agent)",
          "Activity (agent)", FORM_COL, SUBTYPE_COL, FORM_OTHER_COL, ACTIVITY_COL,
          ACTIVITY_OTHER_COL, CHECKED, COMMENT_COL]
DROPDOWNS = {FORM_COL: FORMS, SUBTYPE_COL: TOOL_SUBTYPES, ACTIVITY_COL: ACTIVITIES,
             CHECKED: ("Yes", "No")}
WIDTHS = {"ID": 8, "Title": 40, "Solution name": 20, "Usage context (RQ2)": 18,
          "Form (agent)": 60, "Activity (agent)": 60, COMMENT_COL: 40}
WRAP = {"Title", "Form (agent)", "Activity (agent)", COMMENT_COL}


def _usage_context(origin: str) -> dict:
    """``{id: published RQb value}`` (the validated value when the authors corrected it)."""
    data = json.loads(step_output_path(origin, 4).read_text(encoding="utf-8"))
    out = {}
    for r in data.get("records", []):
        rqb = (r.get("answers") or {}).get("RQb") or {}
        out[r["id"]] = rqb.get("human_value") or rqb.get("value", "")
    return out


def _agent_cell(answer: dict, subtype: bool) -> str:
    lines = [f"value: {answer.get('value', '')}"]
    if subtype and answer.get("tool_subtype"):
        lines.append(f"tool subtype: {answer['tool_subtype']}")
    if answer.get("other"):
        lines.append(f"other: {answer['other']}")
    lines += [f"note: {answer.get('note', '')}", f"evidence: {answer.get('evidence', '')}"]
    return "\n".join(lines)


def export_review_sheet() -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    path = review_dir() / REVIEW_SHEET
    if path.exists():  # never overwrite a sheet that may already hold the reviewer's work
        raise RuntimeError(f"{path} already exists; delete it explicitly to export it again")
    rows = []
    for origin in ORIGINS:
        data, _ = _load(origin)
        context = _usage_context(origin)
        for r in data.get("records", []):
            rq2 = r.get("rq2") or {}
            form, activity = rq2.get("form") or {}, rq2.get("activity") or {}
            rows.append([
                r["id"], r["data"].get("title", ""), r["data"].get("solution_name", ""),
                context.get(r["id"], ""), _agent_cell(form, True), _agent_cell(activity, False),
                form.get("value", ""), form.get("tool_subtype", ""), form.get("other", ""),
                activity.get("value", ""), activity.get("other", ""), "", "",
            ])
    pending = [row[0] for row in rows if not row[6]]
    if pending:
        raise RuntimeError(f"step 8 has not answered {len(pending)} studies yet: {', '.join(pending)}")

    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "B2"
    for row in rows:
        ws.append(row)
    wrap = Alignment(wrap_text=True, vertical="top")
    for col, name in enumerate(HEADER, start=1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = WIDTHS.get(name, 16)
        if name in WRAP:
            for cell in ws[letter][1:]:
                cell.alignment = wrap
    for name, values in DROPDOWNS.items():
        letter = get_column_letter(HEADER.index(name) + 1)
        validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"',
                                    allow_blank=True)
        ws.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{len(rows) + 1}")
    # The reviewer validates against the same definitions the agent used: the prompt, verbatim.
    criteria = wb.create_sheet("criteria")
    criteria.column_dimensions["A"].width = 110
    for line in PROMPT.read_text(encoding="utf-8").splitlines():
        criteria.append([line])
    ensure_parent(path)
    wb.save(path)
    print(f"[rq2] {REVIEW_SHEET}: {len(rows)} rows -> {path}")
    return {"rows": len(rows), "path": str(path)}


# --- Step 10: apply the validated values ---

def _cell(value) -> str:
    return "" if value is None else str(value).strip()


def apply_review() -> dict:
    from openpyxl import load_workbook

    path = review_dir() / REVIEW_SHEET
    ws = load_workbook(path, read_only=True, data_only=True)["review"]
    rows = list(ws.iter_rows(values_only=True))
    col = {name: i for i, name in enumerate(rows[0])}
    loaded = {origin: _load(origin) for origin in ORIGINS}

    stats = {"checked": 0, "unchecked": [], "form_changed": 0, "subtype_changed": 0,
             "activity_changed": 0, "invalid": [], "unknown": []}
    for row in rows[1:]:
        # A sheet round-tripped through Google Sheets comes back with trailing empty rows, which
        # read-only openpyxl yields as short tuples; pad them so every column lookup is safe.
        row = tuple(row) + (None,) * (len(col) - len(row))
        rid = _cell(row[col["ID"]])
        if not rid:
            continue
        origin = ORIGIN_BY_PREFIX.get(rid[:2])
        record = loaded[origin][1].get(rid) if origin else None
        if record is None or not record.get("rq2"):
            stats["unknown"].append(rid)
            continue
        if _cell(row[col[CHECKED]]).lower() != "yes":
            stats["unchecked"].append(rid)
            continue
        form, activity = record["rq2"]["form"], record["rq2"]["activity"]
        new_form, new_subtype = _cell(row[col[FORM_COL]]), _cell(row[col[SUBTYPE_COL]])
        new_activity = _cell(row[col[ACTIVITY_COL]])
        if (new_form not in FORMS or new_activity not in ACTIVITIES
                or (new_subtype and new_subtype not in TOOL_SUBTYPES)):
            stats["invalid"].append(rid)
            continue
        # human_* keys are written only where the reviewer changed the model's answer, so the
        # model's own value stays intact for comparison (same convention as step 4).
        for answer, key, new, counter in ((form, "value", new_form, "form_changed"),
                                          (form, "tool_subtype", new_subtype, "subtype_changed"),
                                          (activity, "value", new_activity, "activity_changed")):
            human_key = "human_value" if key == "value" else f"human_{key}"
            answer.pop(human_key, None)
            if new != answer.get(key, ""):
                answer[human_key] = new
                stats[counter] += 1
        for answer, other_col in ((form, FORM_OTHER_COL), (activity, ACTIVITY_OTHER_COL)):
            answer.pop("human_other", None)
            new_other = _cell(row[col[other_col]])
            if new_other != answer.get("other", ""):
                answer["human_other"] = new_other
        record["validated"] = True
        record["review_comment"] = _cell(row[col[COMMENT_COL]])
        stats["checked"] += 1

    for origin, (data, by_id) in loaded.items():
        if data:
            _save(origin, list(by_id.values()), data.get("model"))
    print(f"[rq2] applied {stats['checked']} checked rows; corrections: form "
          f"{stats['form_changed']}, tool subtype {stats['subtype_changed']}, activity "
          f"{stats['activity_changed']}")
    for label in ("unchecked", "invalid", "unknown"):
        if stats[label]:
            print(f"[rq2] {label} ({len(stats[label])}): {', '.join(stats[label])}")
    return stats
